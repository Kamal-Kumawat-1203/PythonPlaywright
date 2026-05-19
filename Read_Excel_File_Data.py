import openpyxl
import os
# import os helps us to read file in our system or Operating system

file_path = r"C:\Users\kumaw\Desktop\testdata.xlsx"
# r"..." avoids backslash escape issues in Windows paths
if not os.path.exists(file_path):
    print(f"File not found: {file_path}")
else:
    workbook = openpyxl.load_workbook(file_path)
    if "LoginTests" not in workbook.sheetnames:
        print("Sheet 'LoginTests' not found")
    else:
        sheet = workbook["LoginTests"]
        totalrows = sheet.max_row
        totalcols = sheet.max_column
        print("Total rows are:", totalrows, "and Total columns are:", totalcols)

print(sheet.cell(row=3, column=1).value)

# # Below loop will return data in the next line after reading a cell
# for rows in range(1, totalrows + 1):
#     for cols in range(1, totalcols + 1):
#         print(sheet.cell(row=rows, column=cols).value)

# Below loop will print a row data continuously
for rows in range(1, totalrows + 1):
    for cols in range(1, totalcols + 1):
        print(sheet.cell(row=rows, column=cols).value, end=" ")
        # Here end is used to give space after printing a cell of the file
    print()             # Here print function is to move cursor to the next line
